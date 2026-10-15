"""Render normalized HTTPS findings. No network, MCP, or paid API calls."""
import argparse
import json
import re
from datetime import date
from pathlib import Path

SPECS = {
    "http": ("9.1", "9.1 HTTPS vs HTTP", ["address", "issue_description", "suggestion"]),
    "mixed": ("9.2", "9.2 Mixed Content", ["page_address", "resource_url", "issue_description", "suggestion"]),
    "hostname": ("9.3", "9.3 WWW vs non-WWW", ["test_url", "expected_url", "issue_description", "suggestion"]),
}
DETAIL_SHEET = "9. Implement HTTPS Protocol"
DETAIL_HEADERS = ["Issue Type", "Address", "Related URL", "Findings", "Instructions"]
OVERVIEW_HEADERS = ["Item No.", "Item Name", "Initial Check", "Findings", "Instruction"]

CHECK_NAMES = {
    "9.1": "HTTP vs HTTPS (Insecure Content Detected?) - Screaming Frog Insecure Content",
    "9.2": "Mixed Content",
    "9.3": "WWW vs non-WWW",
}


def report_filename(site_name, audit_date):
    if not isinstance(site_name, str) or not site_name.strip():
        raise ValueError("site name is required")
    if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", audit_date):
        raise ValueError("audit date must use YYYY-MM-DD")
    date.fromisoformat(audit_date)
    # Preserve readable site names; replace only cross-platform unsafe characters.
    name = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", site_name.strip()).strip(" .")
    if not name:
        raise ValueError("site name contains no usable filename characters")
    return f"9.1-9.3_{name}_https_audit_{audit_date}.xlsx"


def normalize(data):
    overview = data.get("overview", [])
    if len(overview) != 3 or {r.get("check") for r in overview} != {"9.1", "9.2", "9.3"}:
        raise ValueError("overview must contain one row each for 9.1, 9.2, 9.3")
    by_check = {r["check"]: r for r in overview}
    rows = {}
    for key, (check, _, fields) in SPECS.items():
        unique = {}
        for row in data.get(key, []):
            required = fields + ["issue"]
            if any(not isinstance(row.get(f), str) or not row[f].strip() for f in required):
                raise ValueError(f"{key}: missing/non-string {required}")
            identity = tuple(row[f] for f in fields[:-2]) + (row["issue"],)
            if identity in unique and unique[identity] != row:
                raise ValueError(f"{key}: conflicting duplicate {identity}")
            unique[identity] = dict(row)
        rows[key] = list(unique.values())
        summary = by_check[check]
        if summary.get("result") in {"Needs Review", "Not Tested"}:
            summary = dict(summary, result="Human Check")
            by_check[check] = summary
        if summary.get("result") not in {"Pass", "Issue", "Human Check"}:
            raise ValueError(f"{check}: invalid result")
        for field in ("findings", "instructions"):
            value = summary.get(field)
            if not isinstance(value, str) or not value.strip() or re.fullmatch(r"[\d\s.,+-]+", value):
                raise ValueError(f"{check}: {field} requires a concise sentence, not a count; update findings.json")
        if rows[key] and summary["result"] == "Pass":
            raise ValueError(f"{check}: findings contradict result")
        if not rows[key] and summary["result"] == "Issue":
            raise ValueError(f"{check}: Issue requires evidence rows")
        if summary["result"] == "Issue" and all(r["issue"].startswith(("review:", "human-check:")) for r in rows[key]):
            raise ValueError(f"{check}: only review rows, use Human Check")
        if summary["result"] == "Human Check" and any(not r["issue"].startswith(("review:", "human-check:")) for r in rows[key]):
            raise ValueError(f"{check}: confirmed defects require Issue result")
        for row in rows[key]:
            if row["issue"].startswith(("review:", "human-check:")) and not row["issue_description"].startswith("Human Check"):
                row["issue_description"] = "Human Check: " + row["issue_description"]
    return [by_check[r["check"]] for r in overview], rows


def incomplete_data(reason):
    return {"overview": [{"check": check, "result": "Human Check",
            "findings": reason, "instructions": "请人工检查此项，或恢复连接后补查。"}
            for check in CHECK_NAMES], "http": [], "mixed": [], "hostname": []}


def prepare_report(data, statuses=()):
    """Keep valid per-check results; expose invalid or interrupted checks for review."""
    result = incomplete_data("本项未完成自动检查。")
    source = data if isinstance(data, dict) else {}
    summaries = source.get("overview", [])
    if not isinstance(summaries, list):
        summaries = []
    for index, (key, (check, *_)) in enumerate(SPECS.items()):
        matches = [r for r in summaries if isinstance(r, dict) and r.get("check") == check]
        candidate = incomplete_data("本项未完成自动检查。")
        if len(matches) == 1:
            candidate["overview"][index] = matches[0]
            candidate[key] = source.get(key, [])
            try:
                overview, rows = normalize(candidate)
                result["overview"][index] = overview[index]
                result[key] = rows[key]
            except (ValueError, TypeError, KeyError, AttributeError):
                result["overview"][index]["findings"] = "本项检查数据不完整，未能确认结果。"
    for status in statuses:
        if status.get("status") not in {"failed", "incomplete"}:
            continue
        affected = status.get("affected_checks", list(CHECK_NAMES))
        for index, (key, (check, *_)) in enumerate(SPECS.items()):
            if check not in affected:
                continue
            summary = result["overview"][index]
            detail = str(status.get("error_kind", "")) + " " + str(status.get("error_detail", ""))
            detail = detail.lower()
            reason = ("请求超时，未完成检查。" if "timeout" in detail or "timed out" in detail
                      else "网站返回 429，请求被限流。" if "429" in detail or "rate_limit" in detail
                      else "连接失败，未完成检查。")
            action = "请人工检查此项，或恢复连接后补查。"
            if summary["result"] == "Issue":
                fields = SPECS[key][2]
                row = {f: "未取得" for f in fields}
                row.update(issue="human-check:interrupted", issue_description=reason, suggestion=action)
                result[key].append(row)
                summary["findings"] += " " + reason
                summary["instructions"] += " " + action
            else:
                result["overview"][index] = dict(check=check, result="Human Check", findings=reason, instructions=action)
    return result


def read_status(directory):
    path = Path(directory) / "run-status.json"
    if not path.exists():
        return {}
    try:
        return json.loads(path.read_text(encoding="utf-8-sig"))
    except (OSError, ValueError):
        return {"status": "incomplete"}


def build(data, output):
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    output = Path(output)
    data = prepare_report(data, [read_status(output.parent)])
    overview, rows = normalize(data)
    if output.exists():
        raise FileExistsError(f"Refusing to overwrite {output}; select a new path")
    output.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Checklist"
    ws.append(OVERVIEW_HEADERS)
    for key, (check, *_rest) in SPECS.items():
        summary = next(r for r in overview if r["check"] == check)
        has_gaps = any(r["issue"].startswith(("review:", "human-check:")) for r in rows[key])
        flag = {"Pass": "√", "Issue": "X + Human Check" if has_gaps else "X", "Human Check": "Human Check"}[summary["result"]]
        ws.append([check, CHECK_NAMES[check], flag, summary["findings"], summary["instructions"]])
    sheet = wb.create_sheet(DETAIL_SHEET)
    sheet.append(DETAIL_HEADERS)
    detail_values = []
    for key, (_, issue_type, fields) in SPECS.items():
        for row in rows[key]:
            related = row.get("related_url", "") if key == "http" else row[fields[1]]
            if not isinstance(related, str):
                raise ValueError(f"{key}: related URL must be a string")
            values = [issue_type, row[fields[0]], related, row["issue_description"], row["suggestion"]]
            detail_values.append(values)
            sheet.append(values)
    for sheet in wb:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="2F75B5")
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                # Force text, including formula-looking untrusted input.
                if isinstance(cell.value, str):
                    cell.data_type = "s"
                    cell.number_format = "@"
                    if cell.value.startswith(("https://", "http://")):
                        cell.hyperlink = cell.value
                        cell.font = Font(color="0563C1", underline="single")
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for col in sheet.columns:
            label = col[0].value
            width = 60 if label in {"Findings", "Instructions", "Instruction"} else 48
            if label == "Issue Type":
                width = 28
            if label in {"Initial Check", "Item No."}:
                width = 18
            sheet.column_dimensions[col[0].column_letter].width = width
    wb.save(output)
    verified = load_workbook(output)
    expected = ["Checklist", DETAIL_SHEET]
    if verified.sheetnames != expected:
        raise RuntimeError("Workbook sheet verification failed")
    if [verified["Checklist"].cell(i, 1).value for i in range(2, 5)] != list(CHECK_NAMES):
        raise RuntimeError("Workbook check name verification failed")
    if [verified["Checklist"].cell(i, 2).value for i in range(2, 5)] != list(CHECK_NAMES.values()):
        raise RuntimeError("Workbook item name verification failed")
    if [c.value for c in verified["Checklist"][1]] != OVERVIEW_HEADERS:
        raise RuntimeError("Overview header verification failed")
    sheet = verified[DETAIL_SHEET]
    if [c.value for c in sheet[1]] != DETAIL_HEADERS or sheet.max_row - 1 != len(detail_values):
        raise RuntimeError("Workbook detail header/count verification failed")
    for i, source in enumerate(detail_values, 2):
        if [sheet.cell(i, j).value or "" for j in range(1, 6)] != source:
            raise RuntimeError("Workbook value preservation failed")
    verified.close()
    return {"output": str(output.resolve()), "findings": {k: len(v) for k, v in rows.items()}}


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", required=True, type=Path)
    parser.add_argument("--output-dir", required=True, type=Path)
    parser.add_argument("--site-name", required=True)
    parser.add_argument("--date", required=True, help="Audit date YYYY-MM-DD in the user's timezone")
    args = parser.parse_args()
    output = args.output_dir / report_filename(args.site_name, args.date)
    try:
        data = json.loads(args.input.read_text(encoding="utf-8-sig"))
    except (OSError, ValueError) as exc:
        data = incomplete_data("未能读取检查数据，本项需要人工检查。")
        args.output_dir.mkdir(parents=True, exist_ok=True)
        (args.output_dir / "report-error.log").write_text(str(exc), encoding="utf-8")
    if args.input.parent.resolve() != args.output_dir.resolve():
        data = prepare_report(data, [read_status(args.input.parent)])
    print(json.dumps(build(data, output), ensure_ascii=False))
