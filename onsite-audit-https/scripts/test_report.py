"""Offline behavioral checks using synthetic evidence."""
import tempfile
import unittest
from pathlib import Path
from openpyxl import load_workbook
from build_report import build, normalize, report_filename


def fixture():
    return {"overview": [
        {"check": "9.1", "result": "Issue", "findings": "HTTP page returns 200.", "instructions": "Configure HTTPS redirect."},
        {"check": "9.2", "result": "Pass", "findings": "No mixed content found in supplied sample.", "instructions": "No action required."},
        {"check": "9.3", "result": "Human Check", "findings": "Hostname returned 403.", "instructions": "Inspect the blocked page."}],
        "http": [{"address": "http://example.com/a/?x=1&y=2", "issue": "http-200", "issue_description": "HTTP page returns 200", "suggestion": "=untrusted text"}],
        "mixed": [], "hostname": [{"test_url":"https://example.com/", "expected_url":"https://www.example.com/", "issue":"human-check:403", "issue_description":"403 prevents content verification.", "suggestion":"Inspect the blocked page."}]}


class ReportTests(unittest.TestCase):
    def test_report_naming_and_date_validation(self):
        self.assertEqual(report_filename("Example", "2026-10-05"), "Example_https_audit_2026-10-05.xlsx")
        self.assertEqual(report_filename("品牌/香港", "2026-10-05"), "品牌_香港_https_audit_2026-10-05.xlsx")
        with self.assertRaises(ValueError):
            report_filename("Example", "2026-02-30")
        with self.assertRaises(ValueError):
            report_filename("", "2026-10-05")

    def test_preserves_urls_and_text_in_merged_sheet(self):
        data = fixture()
        data["http"].append(dict(data["http"][0]))
        # Keep test artifacts under the caller-selected workspace.
        with tempfile.TemporaryDirectory(dir=Path.cwd()) as directory:
            output = Path(directory) / "audit.xlsx"
            result = build(data, output)
            self.assertEqual(result["findings"]["http"], 1)
            wb = load_workbook(output)
            self.assertEqual(wb.sheetnames, ["Overview", "9. Implement HTTPS Protocol"])
            self.assertEqual(wb.worksheets[1]["B2"].value, data["http"][0]["address"])
            self.assertEqual(wb.worksheets[1]["D2"].value, "HTTP page returns 200")
            self.assertEqual(wb.worksheets[1]["E2"].value, "=untrusted text")
            self.assertEqual(wb.worksheets[1]["E2"].data_type, "s")
            self.assertEqual(wb["Overview"]["B4"].value, "Human Check")
            self.assertEqual(wb["Overview"]["B2"].value, "X")
            self.assertEqual([wb["Overview"].cell(i, 1).value for i in range(2, 5)], [
                "9.1 HTTP vs HTTPS (Insecure Content Detected?) - Screaming Frog Insecure Content",
                "9.2 Mixed Content", "9.3 WWW vs non-WWW"])
            wb.close()
            with self.assertRaises(FileExistsError):
                build(data, output)

    def test_rejects_false_pass_and_missing_resource(self):
        data = fixture()
        data["overview"][0]["result"] = "Pass"
        with self.assertRaises(ValueError):
            normalize(data)
        data = fixture()
        data["mixed"] = [{"page_address": "https://example.com", "issue": "mixed", "suggestion": "Fix resource"}]
        with self.assertRaises(ValueError):
            normalize(data)

    def test_retains_shared_resource_on_different_pages(self):
        data = fixture()
        data["overview"][1]["result"] = "Issue"
        data["mixed"] = [{"page_address": page, "resource_url": "http://example.com/logo.png", "issue": "mixed", "issue_description": "HTTPS page loads an HTTP image", "suggestion": "Update image reference"} for page in ["https://example.com/a", "https://example.com/b"]]
        _, rows = normalize(data)
        self.assertEqual(len(rows["mixed"]), 2)

    def test_tick_is_rendered_without_treating_unknown_as_pass(self):
        data = fixture()
        data["overview"][1] = {"check": "9.2", "result": "Pass", "findings": "No mixed content found.", "instructions": "No action required."}
        with tempfile.TemporaryDirectory(dir=Path.cwd()) as directory:
            output = Path(directory) / "audit.xlsx"
            build(data, output)
            wb = load_workbook(output)
            self.assertEqual(wb["Overview"]["B3"].value, "√")
            self.assertEqual(wb["Overview"]["B4"].value, "Human Check")
            self.assertEqual(wb["Overview"]["B2"].value, "X")
            wb.close()

    def test_partial_defect_keeps_x_and_exports_human_check(self):
        data = fixture()
        data['http'].append({'address':'https://example.com/missing', 'issue':'human-check:response',
                             'issue_description':'Final response unverified; HTTP 403.', 'suggestion':'Manually inspect redirect and destination.'})
        with tempfile.TemporaryDirectory(dir=Path.cwd()) as directory:
            build(data,Path(directory)/'audit.xlsx')
            wb=load_workbook(Path(directory)/'audit.xlsx')
            self.assertEqual(wb['Overview']['B2'].value,'X + Human Check')
            self.assertTrue(wb['9. Implement HTTPS Protocol']['D3'].value.startswith('Human Check'))
            self.assertEqual(wb['9. Implement HTTPS Protocol']['E3'].value,'Manually inspect redirect and destination.')
            self.assertTrue(wb['9. Implement HTTPS Protocol']['D4'].value.startswith('Human Check'))
            wb.close()

    def test_sample_zero_no_review_rows(self):
        data = {"overview": [{"check": check, "result": "Pass",
                "findings": "No issues found in the supplied sample.",
                "instructions": "No action required."} for check in ["9.1", "9.2", "9.3"]],
                "http": [], "mixed": [], "hostname": []}
        with tempfile.TemporaryDirectory(dir=Path.cwd()) as directory:
            output = Path(directory) / "audit.xlsx"
            build(data, output)
            wb = load_workbook(output)
            self.assertEqual(wb.sheetnames, ["Overview", "9. Implement HTTPS Protocol"])
            self.assertEqual([c.value for c in wb['Overview'][1]], ["Check", "Flag", "Findings", "Instructions"])
            self.assertEqual([c.value for c in wb.worksheets[1][1]], ["Issue Type", "Address", "Related URL", "Findings", "Instructions"])
            self.assertEqual(wb.worksheets[1].max_row, 1)
            self.assertTrue(all(wb['Overview'].cell(i, 2).value == "√" for i in range(2, 5)))
            self.assertEqual(wb['Overview']['C2'].value, data['overview'][0]['findings'])
            wb.close()

    def test_no_invented_human_check_or_numeric_findings(self):
        data = fixture()
        data['hostname'] = []
        with self.assertRaises(ValueError): normalize(data)
        for value in [1, "1", ""]:
            data = fixture()
            data['overview'][0]['findings'] = value
            with self.assertRaises(ValueError): normalize(data)

    def test_issue_types_and_mixed_resource_survive_merge(self):
        data = fixture()
        data['overview'][1]['result'] = 'Issue'
        data['mixed'] = [{"page_address":"https://example.com/a", "resource_url":"http://example.com/a.png",
            "issue":"mixed", "issue_description":"HTTP image on HTTPS page.", "suggestion":"Update image URL."}]
        with tempfile.TemporaryDirectory(dir=Path.cwd()) as directory:
            output = Path(directory) / 'audit.xlsx'
            build(data, output)
            wb = load_workbook(output)
            sheet = wb.worksheets[1]
            self.assertEqual([sheet.cell(i, 1).value for i in range(2, 5)],
                ['9.1 HTTPS vs HTTP', '9.2 Mixed Content', '9.3 WWW vs non-WWW'])
            self.assertEqual(sheet['B3'].value, data['mixed'][0]['page_address'])
            self.assertEqual(sheet['C3'].value, data['mixed'][0]['resource_url'])
            wb.close()

    def test_review_only_rows_cannot_be_confirmed_issue(self):
        data=fixture()
        data['http'][0]['issue']='human-check:timeout'
        with self.assertRaises(ValueError):normalize(data)


if __name__ == "__main__":
    unittest.main()
